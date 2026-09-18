#!/usr/bin/env python3
"""Priority-date density from DOL labour certification disclosure data.

This is the missing signal. The backtest showed the estimate leans optimistic
because it reads how fast the cutoff has moved without knowing how many people
hold each priority date: years that moved quickly did so partly because the
cohorts behind the cutoff were thin. PLAN.md finding 32 and the backtest
results.

A certified PERM's CASE_RECEIVED_DATE **is** the priority date, so counting
certified cases by received month gives the shape of the queue.

Two things this data does better than the plan assumed:

- Files up to FY2019 carry `FW_INFO_BIRTH_COUNTRY`, not only citizenship. The
  plan recorded citizenship-as-proxy as a permanent caveat; for those years it
  is not needed, and those are the years covering the 2013-2017 priority dates
  where the India queue actually sits. From FY2020 the field was dropped and
  only citizenship remains, so the caveat applies to newer cohorts only. Each
  source year records which field it used.
- It carries `JOB_INFO_EDUCATION`, the job's minimum requirement, which
  separates likely EB-2 from likely EB-3 demand. Still a proxy: the real
  category is on the I-140, not the PERM.

What it still cannot see: EB-1 and national interest waiver cases, which never
file a PERM, and whether a certified PERM ever became an I-140.

Files are large (55-90 MB each). Each year is downloaded, streamed, aggregated
and the raw file discarded. Already-aggregated years are skipped, so the run is
resumable.

Usage:
    python build_perm_density.py                 # every known year
    python build_perm_density.py --years 2015 2016
"""

from __future__ import annotations

import argparse
import json
import sys
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import openpyxl  # noqa: E402
import requests  # noqa: E402

from gceta.fetch import USER_AGENT  # noqa: E402

DATA_DIR = Path(__file__).resolve().parents[1] / "data"
CACHE = Path(__file__).resolve().parent / "cache"
OUT = DATA_DIR / "perm-density.json"

BASE = "https://www.dol.gov/sites/dolgov/files/ETA/oflc/pdfs/"

# Naming is inconsistent across years, and guessing at names wasted a lot of
# time. The authoritative list is the OFLC performance page, which links every
# published disclosure file; scraping its hrefs for /PERM.*\.xlsx?/ found both
# of the years previously recorded as missing, under names no pattern guessed:
# FY2014 as PERM_FY14_Q4.xlsx and FY2018 as PERM_Disclosure_Data_FY2018_EOY.xlsx.
#
# THE RECORD CANNOT REACH BACK BEFORE FY2015, and not because of naming.
# Files for FY2008 through FY2014 are published and downloadable, but they carry
# 27 columns: CASE_NO, DECISION_DATE, CASE_STATUS, employer and wage fields,
# COUNTRY_OF_CITIZENSHIP, CLASS_OF_ADMISSION. There is no CASE_RECEIVED_DATE.
# The priority date IS the receipt date, so a decision date cannot substitute
# without inventing a processing-time distribution for exactly the years the
# estimate is most sensitive to. Those years are excluded on purpose. From
# FY2015 the file widens to 125 columns and carries receipt date, the job's
# education requirement, and birth country.
#
# The practical floor this sets: a PERM decided in FY2015 was typically received
# in 2013 or 2014, so priority-date coverage begins around 2012 and is thin
# until 2013. India's employment cutoffs currently sit at the edge of that,
# which is why the queue count works going forward from today but not for
# historical backtest months when cutoffs sat in the 2000s.
FILES = {
    2015: "PERM_Disclosure_Data_FY15_Q4.xlsx",
    2016: "PERM_Disclosure_Data_FY16.xlsx",
    2017: "PERM_Disclosure_Data_FY17.xlsx",
    # Found on the OFLC index. The file previously tried, PERM_FY2018.xlsx,
    # holds only Q2; this one holds 119,776 rows across the full year and fills
    # what looked like a real dip in 2017-2018 priority dates but was a gap.
    2018: "PERM_Disclosure_Data_FY2018_EOY.xlsx",
    2019: "PERM_Disclosure_Data_FY2019.xlsx",
    2020: "PERM_Disclosure_Data_FY2020_Q4.xlsx",
    2021: "PERM_Disclosure_Data_FY2021_Q4.xlsx",
    2022: "PERM_Disclosure_Data_FY2022_Q4.xlsx",
    2023: "PERM_Disclosure_Data_FY2023_Q4.xlsx",
    2024: "PERM_Disclosure_Data_FY2024_Q4.xlsx",
    2025: "PERM_Disclosure_Data_FY2025_Q4.xlsx",
}
# Decision-date-only schema, so no priority date can be derived. Not a gap that
# further searching will close.
NO_RECEIPT_DATE = [2008, 2009, 2010, 2011, 2012, 2013, 2014]
KNOWN_MISSING = []

COLUMN_FOR_COUNTRY = {
    "INDIA": "IN",
    "CHINA": "CN",
    "CHINA - HONG KONG": "ROW",
    "CHINA-HONG KONG": "ROW",
    "MEXICO": "MX",
    "PHILIPPINES": "PH",
}

# Only a certified case can support an I-140, so only certified cases are queue.
COUNTED_STATUSES = {"certified", "certified-expired"}


def bucket_education(value: str | None) -> str:
    text = (value or "").strip().lower()
    if text.startswith("master") or "doctor" in text or text.startswith("phd"):
        return "advanced"
    if text.startswith("bachelor"):
        return "bachelors"
    return "other"


def download(year: int) -> Path:
    name = FILES[year]
    path = CACHE / name
    if path.exists() and path.stat().st_size > 1_000_000:
        return path
    CACHE.mkdir(parents=True, exist_ok=True)
    response = requests.get(
        BASE + name, headers={"User-Agent": USER_AGENT}, timeout=900, stream=True
    )
    response.raise_for_status()
    with open(path, "wb") as handle:
        for chunk in response.iter_content(1 << 20):
            handle.write(chunk)
    return path


def aggregate(path: Path) -> dict:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    counts: dict[str, dict[str, dict[str, int]]] = defaultdict(
        lambda: defaultdict(lambda: defaultdict(int))
    )
    rows_seen = 0
    counted = 0
    bases: set[str] = set()

    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        stream = sheet.iter_rows(values_only=True)
        try:
            header = [str(h).strip().upper() if h else "" for h in next(stream)]
        except StopIteration:
            continue
        index = {name: i for i, name in enumerate(header)}

        # The schema was renamed for FY2020: CASE_RECEIVED_DATE became
        # RECEIVED_DATE, JOB_INFO_EDUCATION became MINIMUM_EDUCATION, and
        # FW_INFO_BIRTH_COUNTRY was dropped. Silently skipping those files
        # produced four years of zero rows; the coverage report caught it.
        def first_present(*names: str) -> str | None:
            for name in names:
                if name in index:
                    return name
            return None

        received_key = first_present("CASE_RECEIVED_DATE", "RECEIVED_DATE")
        if "CASE_STATUS" not in index or received_key is None:
            continue
        country_key = first_present("FW_INFO_BIRTH_COUNTRY", "COUNTRY_OF_CITIZENSHIP")
        education_key = first_present("JOB_INFO_EDUCATION", "MINIMUM_EDUCATION")
        country_basis = (
            "birth" if country_key == "FW_INFO_BIRTH_COUNTRY" else "citizenship"
        )
        bases.add(country_basis)

        for row in stream:
            rows_seen += 1
            status = row[index["CASE_STATUS"]]
            if not status or str(status).strip().lower() not in COUNTED_STATUSES:
                continue
            received = row[index[received_key]]
            if not isinstance(received, datetime):
                continue
            country = str(row[index.get(country_key, -1)] or "").strip().upper()
            column = COLUMN_FOR_COUNTRY.get(country, "ROW")
            month = f"{received.year:04d}-{received.month:02d}"
            education = bucket_education(
                row[index[education_key]] if education_key else None
            )
            counts[column][month]["total"] += 1
            counts[column][month][education] += 1
            counted += 1

    workbook.close()
    return {
        "rows_seen": rows_seen,
        "certified": counted,
        "months": sorted({m for months in counts.values() for m in months}),
        "country_basis": sorted(bases),
        "counts": {c: {m: dict(v) for m, v in months.items()} for c, months in counts.items()},
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--years", type=int, nargs="*", default=sorted(FILES))
    parser.add_argument("--keep-raw", action="store_true")
    args = parser.parse_args()

    existing = json.loads(OUT.read_text()) if OUT.exists() else {"years": {}, "density": {}}
    density: dict[str, dict[str, dict[str, int]]] = {
        c: {m: dict(v) for m, v in months.items()}
        for c, months in existing.get("density", {}).items()
    }
    done = existing.get("years", {})

    for year in args.years:
        if str(year) in done:
            print(f"  FY{year}  already aggregated ({done[str(year)]['certified']:,} certified)")
            continue
        if year not in FILES:
            print(f"  FY{year}  no published file")
            continue
        print(f"  FY{year}  downloading…", flush=True)
        path = download(year)
        print(f"  FY{year}  parsing {path.stat().st_size/1048576:.0f} MB…", flush=True)
        result = aggregate(path)
        for column, months in result["counts"].items():
            target = density.setdefault(column, {})
            for month, values in months.items():
                bucket = target.setdefault(month, {})
                for key, value in values.items():
                    bucket[key] = bucket.get(key, 0) + value
        months = result["months"]
        done[str(year)] = {
            "certified": result["certified"],
            "rows": result["rows_seen"],
            "received_months": len(months),
            "span": [months[0], months[-1]] if months else None,
            "country_basis": result["country_basis"],
        }
        months = result["months"]
        span = f"{months[0]}..{months[-1]}" if months else "none"
        print(
            f"  FY{year}  {result['certified']:,} certified of {result['rows_seen']:,} rows"
            f"  ·  {len(months)} received-months  ({span})"
        )
        if not args.keep_raw:
            path.unlink(missing_ok=True)

        OUT.write_text(json.dumps({
            "schema_version": 1,
            "generated_at": datetime.now(timezone.utc).isoformat(),
            "source": "DOL OFLC PERM disclosure data",
            "note": (
                "Certified labour certifications counted by the month DOL received "
                "them, which is the priority date. Birth country where published, "
                "citizenship otherwise. Education buckets proxy the EB-2 / EB-3 "
                "split; the real category is on the I-140, not the PERM. Excludes "
                "EB-1 and national interest waiver cases, which file no PERM."
            ),
            "years_missing": KNOWN_MISSING,
            "years_without_receipt_date": NO_RECEIPT_DATE,
            "years": done,
            "density": density,
        }, separators=(",", ":")))

    total = sum(v["certified"] for v in done.values())
    print(f"\nyears aggregated : {len(done)}  (missing: {KNOWN_MISSING})")
    print(f"no receipt date  : FY{NO_RECEIPT_DATE[0]}-FY{NO_RECEIPT_DATE[-1]} publish decision date only")
    print(f"certified cases  : {total:,}")
    print(f"written          : {OUT} ({OUT.stat().st_size/1024:.0f} KB)")

    # Coverage, measured rather than assumed. A fiscal year file that yields far
    # fewer than twelve received-months is scoped to part of the year whatever
    # its name says, and silently undercounts the queue.
    thin = [
        (y, v) for y, v in sorted(done.items())
        if v.get("received_months") is not None and v["received_months"] < 10
    ]
    if thin:
        print("\nWARNING, these source years look partial:")
        for y, v in thin:
            print(f"   FY{y}: only {v['received_months']} received-months {v.get('span')}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
